"""Collector, lifecycle, and CLI tests. All network mocked."""

import json
from datetime import date

import httpx
import pytest
import respx
from typer.testing import CliRunner

from jobscout.collect import COLLECTORS, CollectorError, RawPosting
from jobscout.collect.runner import run_collection
from jobscout.collect.text import html_to_text, remote_from, salary_from_text, workday_posted
from jobscout.http import PoliteClient
from jobscout.jobs import known_ids, save_postings
from tests.conftest import load_fixture


@pytest.fixture
def http():
    with PoliteClient(min_interval=0, backoff=0, retries=0) as c:
        yield c


# ---------------------------------------------------------------------------
# Text helpers
# ---------------------------------------------------------------------------


def test_html_to_text_handles_escaped_greenhouse_html():
    t = html_to_text("&lt;p&gt;Hello &amp;amp; welcome&lt;/p&gt;&lt;ul&gt;&lt;li&gt;One&lt;/li&gt;&lt;/ul&gt;")
    assert "Hello & welcome" in t and "- One" in t and "<" not in t


@pytest.mark.parametrize("text, expected", [
    ("Pay range: $140,000 - $175,000 USD", (140000, 175000)),
    ("between $120k and ... $120k-$150k base", (120000, 150000)),
    ("$115,000 to $150,000 annually", (115000, 150000)),
    ("$25 - $40 per hour", None),
    ("no pay listed", None),
])
def test_salary_from_text(text, expected):
    assert salary_from_text(text) == expected


@pytest.mark.parametrize("hints, expected", [
    (("Remote - United States",), True),
    (("Hybrid - Chicago",), False),
    (("Milwaukee, WI",), None),
    ((None, "Remote"), True),
])
def test_remote_from(hints, expected):
    assert remote_from(*hints) is expected


def test_workday_posted():
    today = date(2026, 10, 11)
    assert workday_posted("Posted Today", today) == "2026-10-11"
    assert workday_posted("Posted Yesterday", today) == "2026-10-10"
    assert workday_posted("Posted 30+ Days Ago", today) == "2026-09-11"


# ---------------------------------------------------------------------------
# Collectors
# ---------------------------------------------------------------------------


@respx.mock
def test_greenhouse(http):
    respx.get("https://boards-api.greenhouse.io/v1/boards/ncinoinc/jobs").respond(
        200, json=load_fixture("collect/greenhouse_jobs.json"))
    [a, b] = COLLECTORS["greenhouse"](http, "ncinoinc")
    assert (a.external_id, a.title, a.remote, a.department) == (
        "7001", "Senior Implementation Manager", True, "Professional Services")
    assert (a.comp_min, a.comp_max, a.comp_source) == (140000, 175000, "text")
    assert a.posted_at == "2026-10-01"
    assert "Own go-lives" in a.description
    assert b.remote is None and b.comp_min is None


@respx.mock
def test_lever(http):
    respx.get("https://api.lever.co/v0/postings/truv").respond(
        200, json=load_fixture("collect/lever_postings.json"))
    [a, b] = COLLECTORS["lever"](http, "truv")
    assert (a.title, a.remote, a.comp_min, a.comp_max, a.comp_source) == (
        "Solutions Architect", True, 150000, 185000, "ats")
    assert "Lead integrations" in a.description and "What you'll do" in a.description
    assert b.remote is False


@respx.mock
def test_ashby_skips_unlisted_and_reads_salary(http):
    respx.get("https://api.ashbyhq.com/posting-api/job-board/evenup").respond(
        200, json=load_fixture("collect/ashby_board.json"))
    [a] = COLLECTORS["ashby"](http, "evenup")
    assert (a.title, a.remote, a.comp_min, a.comp_max) == ("Implementation Lead", True, 130000, 160000)
    assert a.posted_at == "2026-10-05"


@respx.mock
def test_workable(http):
    respx.get("https://apply.workable.com/api/v1/widget/accounts/partner-one-capital").respond(
        200, json=load_fixture("collect/workable_account.json"))
    [a] = COLLECTORS["workable"](http, "partner-one-capital")
    assert (a.external_id, a.remote, a.location) == ("AB12CD", True, "United States")
    assert "LLMs" in a.description


WD = "https://fis.wd5.myworkdayjobs.com/wday/cxs/fis/SearchJobs"


def _mock_workday():
    pages = [load_fixture("collect/workday_page1.json"), load_fixture("collect/workday_page2.json")]
    search = respx.post(f"{WD}/jobs").mock(side_effect=[httpx.Response(200, json=p) for p in pages])
    remote_detail = load_fixture("collect/workday_detail.json")

    def detail_for(request):
        if "JR0301" in request.url.path:
            return httpx.Response(200, json=remote_detail)
        return httpx.Response(200, json={"jobPostingInfo": {"jobDescription": "<p>Onsite role.</p>",
                                                            "remoteType": "On-site"}})

    detail = respx.get(url__startswith=f"{WD}/job/").mock(side_effect=detail_for)
    return search, detail


@respx.mock
def test_workday_pages_and_fetches_details(http):
    search, detail = _mock_workday()
    posts = COLLECTORS["workday"](http, "fis|wd5|SearchJobs")
    assert search.call_count == 2 and detail.call_count == 3
    assert [p.title for p in posts] == [
        "Consultant Lead, Conversions", "Senior Implementation Manager", "Product Owner, Digital Banking"]
    a = posts[0]
    assert a.remote is True and (a.comp_min, a.comp_max) == (115000, 150000)
    assert a.url.startswith("https://fis.wd5.myworkdayjobs.com/SearchJobs/job/")


@respx.mock
def test_workday_skips_details_for_known_postings(http):
    search, detail = _mock_workday()
    known = frozenset({"/job/Remote-USA/Consultant-Lead--Conversions_JR0301",
                       "/job/Milwaukee-WI/Senior-Implementation-Manager_JR0302"})
    posts = COLLECTORS["workday"](http, "fis|wd5|SearchJobs", known=known)
    assert detail.call_count == 1 and len(posts) == 3
    assert posts[0].description is None  # known: list fields only


@respx.mock
def test_collector_errors_on_bad_answer(http):
    respx.get("https://boards-api.greenhouse.io/v1/boards/gone/jobs").respond(404)
    with pytest.raises(CollectorError):
        COLLECTORS["greenhouse"](http, "gone")


# ---------------------------------------------------------------------------
# Lifecycle
# ---------------------------------------------------------------------------


def _company(db, name="Q2", ats="greenhouse", token="q2", **extra):
    cols = {"name": name, "tier": "A", "ats": ats, "ats_token": token,
            "detect_status": "found", "detect_confidence": "confirmed", **extra}
    db.execute(f"INSERT INTO companies ({', '.join(cols)}) VALUES ({', '.join('?' * len(cols))})",
               list(cols.values()))
    return db.execute("SELECT * FROM companies WHERE name = ?", (name,)).fetchone()


def _post(i, title="Implementation Manager", desc="Do the work.", **kw):
    return RawPosting(source="greenhouse", external_id=str(i), title=title, description=desc,
                      payload={"id": i}, **kw)


def test_lifecycle_new_unchanged_updated_closed_reopened(db):
    c = _company(db)
    r = save_postings(db, c["id"], [_post(1), _post(2, "Delivery Manager")])
    assert len(r.new) == 2 and r.open_count == 2

    r = save_postings(db, c["id"], [_post(1), _post(2, "Delivery Manager")])
    assert r.unchanged == 2 and not (r.new or r.updated or r.closed)

    r = save_postings(db, c["id"], [_post(1, desc="Do the work, remotely.")])
    assert r.updated == ["Implementation Manager"] and r.closed == ["Delivery Manager"]

    r = save_postings(db, c["id"], [_post(1, desc="Do the work, remotely."), _post(2, "Delivery Manager")])
    assert r.reopened == ["Delivery Manager"]
    assert db.execute("SELECT COUNT(*) FROM jobs WHERE closed_at IS NULL").fetchone()[0] == 2


def test_missing_description_never_wipes_a_stored_one(db):
    c = _company(db)
    save_postings(db, c["id"], [_post(1, desc="Full description")])
    r = save_postings(db, c["id"], [_post(1, desc=None)])
    assert r.unchanged == 1
    assert db.execute("SELECT description FROM jobs").fetchone()[0] == "Full description"


def test_first_description_is_not_an_update(db):
    c = _company(db)
    save_postings(db, c["id"], [_post(1, desc=None)])
    r = save_postings(db, c["id"], [_post(1, desc="Now with details")])
    assert r.unchanged == 1 and not r.updated


def test_raw_payload_stored_for_new_and_changed_only(db):
    c = _company(db)
    save_postings(db, c["id"], [_post(1)], run_id=None)
    save_postings(db, c["id"], [_post(1)])
    save_postings(db, c["id"], [_post(1, title="Senior Implementation Manager")])
    assert db.execute("SELECT COUNT(*) FROM raw_postings").fetchone()[0] == 2


def test_known_ids_only_counts_open_postings_with_descriptions(db):
    c = _company(db, ats="workday", token="fis|wd5|SearchJobs")
    wd = lambda i, d: RawPosting(source="workday", external_id=i, title="x", description=d)  # noqa: E731
    save_postings(db, c["id"], [wd("/a", "desc"), wd("/b", None)])
    assert known_ids(db, c["id"], "workday") == frozenset({"/a"})


# ---------------------------------------------------------------------------
# Runner
# ---------------------------------------------------------------------------


@respx.mock
def test_failed_board_does_not_close_jobs(db, http):
    c = _company(db)
    save_postings(db, c["id"], [_post(1), _post(2)])
    respx.get("https://boards-api.greenhouse.io/v1/boards/q2/jobs").respond(503)
    [r] = run_collection(db, http, [db.execute("SELECT * FROM companies").fetchone()])
    assert r.status == "error"
    assert db.execute("SELECT COUNT(*) FROM jobs WHERE closed_at IS NULL").fetchone()[0] == 2
    row = db.execute("SELECT collect_status, open_jobs FROM companies").fetchone()
    assert row["collect_status"] == "error"


def test_runner_skips_unsupported_and_review(db, http):
    _company(db, "Fiserv", ats="custom", token="https://careers.fiserv.com", detect_status="override")
    _company(db, "Unit", ats="ashby", token="unit", detect_confidence="review")
    _company(db, "Polly", ats=None, token=None, detect_status="not_found")
    rows = db.execute("SELECT * FROM companies ORDER BY name").fetchall()
    results = {r.company: r.status for r in run_collection(db, http, rows)}
    assert results == {"Fiserv": "unsupported", "Unit": "review", "Polly": "unsupported"}


# ---------------------------------------------------------------------------
# CLI end to end
# ---------------------------------------------------------------------------


def test_cli_collect_and_jobs(workdir, monkeypatch):
    import functools

    import jobscout.http as jhttp
    from jobscout.cli import app
    from jobscout.store import connect

    monkeypatch.setattr(jhttp, "PoliteClient",
                        functools.partial(jhttp.PoliteClient, min_interval=0, backoff=0))
    runner = CliRunner()
    assert runner.invoke(app, ["init"]).exit_code == 0
    conn = connect(workdir / "data" / "test.db")
    _company(conn, "nCino", token="ncinoinc")
    _company(conn, "FIS", ats="workday", token="fis|wd5|SearchJobs", detect_status="override")
    _company(conn, "Fiserv", ats="custom", token="https://careers.fiserv.com", detect_status="override")
    conn.close()

    with respx.mock:
        respx.get("https://boards-api.greenhouse.io/v1/boards/ncinoinc/jobs").respond(
            200, json=load_fixture("collect/greenhouse_jobs.json"))
        _mock_workday()
        result = runner.invoke(app, ["collect"])
    assert result.exit_code == 0, result.output
    assert "2 boards read, 5 open postings, 5 new, 0 closed" in result.output
    assert "1 not readable yet" in result.output
    assert "Gates: 3 pass, 2 filtered out" in result.output

    result = runner.invoke(app, ["jobs", "--new"])
    assert result.exit_code == 0, result.output
    assert "Senior Implementation Manager" in result.output
    assert "Software Engineer II" not in result.output       # not in include keywords
    assert "$140k-$175k" in result.output

    result = runner.invoke(app, ["jobs", "--remote", "--company", "FIS"])
    assert "Consultant Lead, Conversions" in result.output
    assert "1 postings" in result.output

    # The Jacksonville onsite Product Owner is gated out of the default view, visible with --all
    result = runner.invoke(app, ["jobs", "--company", "FIS"])
    assert "Product Owner, Digital Banking" not in result.output
    result = runner.invoke(app, ["jobs", "--company", "FIS", "--all"])
    assert "Product Owner, Digital Banking" in result.output

    result = runner.invoke(app, ["runs"])
    assert "collect" in result.output and json.dumps({}) not in result.output


def test_cli_export_writes_xlsx(workdir, monkeypatch):
    import functools

    from openpyxl import load_workbook

    import jobscout.http as jhttp
    from jobscout.cli import app
    from jobscout.store import connect

    monkeypatch.setattr(jhttp, "PoliteClient",
                        functools.partial(jhttp.PoliteClient, min_interval=0, backoff=0))
    runner = CliRunner()
    runner.invoke(app, ["init"])
    conn = connect(workdir / "data" / "test.db")
    _company(conn, "nCino", token="ncinoinc")
    conn.close()
    with respx.mock:
        respx.get("https://boards-api.greenhouse.io/v1/boards/ncinoinc/jobs").respond(
            200, json=load_fixture("collect/greenhouse_jobs.json"))
        runner.invoke(app, ["collect"])

    out = workdir / "out.xlsx"
    result = runner.invoke(app, ["export", "--out", str(out), "--no-open"])
    assert result.exit_code == 0, result.output
    ws = load_workbook(out)["Jobs"]
    assert ws.max_row == 2 and ws.cell(2, 4).value == "Senior Implementation Manager"

    result = runner.invoke(app, ["export", "--all", "--out", str(out), "--no-open"])
    assert result.exit_code == 0, result.output
    wb = load_workbook(out)
    ws = wb["Jobs"]
    assert ws["D1"].value == "Title" and ws["P1"].value == "Gate" and ws.max_row == 3
    titles = {ws.cell(r, 4).value: r for r in range(2, 4)}
    r = titles["Senior Implementation Manager"]
    assert ws.cell(r, 7).value == 140000 and ws.cell(r, 12).value == "Yes" and ws.cell(r, 16).value == "pass"
    assert ws.cell(r, 4).hyperlink.target.startswith("https://job-boards.greenhouse.io/")
    assert ws.cell(titles["Software Engineer II"], 16).value == "lane:no-match"
    assert wb["By company"]["C2"].value.startswith("=COUNTIF(")
