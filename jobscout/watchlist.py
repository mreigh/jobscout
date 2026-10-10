"""Watchlist: the curated list of target companies.

Source of truth is a spreadsheet you edit by hand: either the .xlsx (reads the
"Companies" sheet) or a .csv with the same headers. Your real one lives in
profile/ (gitignored); samples/watchlist.sample.csv is a scrubbed demo version.

Importing upserts by company name. Companies removed from the file are marked
inactive, not deleted, so their history survives.
"""

from __future__ import annotations

import csv
import sqlite3
from dataclasses import dataclass, field
from pathlib import Path

from pydantic import BaseModel, ValidationError, field_validator

from jobscout.config import ConfigError

# Spreadsheet header -> field name. Matching ignores case and extra spaces.
COLUMNS = {
    "company": "name",
    "segment": "segment",
    "tier": "tier",
    "why it fits": "why_fit",
    "roles to watch": "roles_to_watch",
    "warm angle": "warm_angle",
    "status / history": "status",
    "status": "status",
    "local": "local_flag",
    "ats": "ats_hint",
    "board url": "board_url",
    "board url (override)": "board_url",
    "notes": "notes",
}
REQUIRED = {"name", "tier"}


class WatchlistEntry(BaseModel):
    name: str
    tier: str
    segment: str | None = None
    why_fit: str | None = None
    roles_to_watch: str | None = None
    warm_angle: str | None = None
    status: str | None = None
    local_flag: str | None = None
    ats_hint: str | None = None
    board_url: str | None = None
    notes: str | None = None

    @field_validator("name")
    @classmethod
    def _name(cls, v: str) -> str:
        v = " ".join(v.split())
        if not v:
            raise ValueError("company name is blank")
        return v

    @field_validator("tier")
    @classmethod
    def _tier(cls, v: str) -> str:
        v = v.strip().upper()
        if v not in {"A", "B", "C"}:
            raise ValueError(f"tier must be A, B, or C (got {v!r})")
        return v

    @field_validator(
        "segment", "why_fit", "roles_to_watch", "warm_angle", "status",
        "local_flag", "ats_hint", "board_url", "notes",
    )
    @classmethod
    def _blank_to_none(cls, v: str | None) -> str | None:
        if v is None:
            return None
        v = str(v).strip()
        return v or None


# ---------------------------------------------------------------------------
# Reading
# ---------------------------------------------------------------------------


def _map_headers(headers: list) -> dict[int, str]:
    mapping = {}
    for i, h in enumerate(headers):
        key = " ".join(str(h or "").lower().split())
        if key in COLUMNS:
            mapping[i] = COLUMNS[key]
    missing = REQUIRED - set(mapping.values())
    if missing:
        raise ConfigError(
            f"Watchlist is missing required column(s): {sorted(missing)}. "
            f"Found headers: {[h for h in headers if h]}"
        )
    return mapping


def _read_rows(path: Path) -> list[list]:
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as f:
            return [row for row in csv.reader(f)]
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        sheet = wb["Companies"] if "Companies" in wb.sheetnames else wb.worksheets[0]
        rows = [list(r) for r in sheet.iter_rows(values_only=True)]
        wb.close()
        return rows
    raise ConfigError(f"Watchlist must be .csv or .xlsx, got {path.name}")


def load_watchlist(path: Path) -> list[WatchlistEntry]:
    if not path.exists():
        raise ConfigError(
            f"Watchlist not found at {path}. Put your spreadsheet there, set "
            "JOBSCOUT_WATCHLIST, or try the demo list with "
            "`--file samples/watchlist.sample.csv`."
        )
    rows = _read_rows(path)
    if not rows:
        raise ConfigError(f"{path} is empty")
    mapping = _map_headers(rows[0])

    entries, errors, seen = [], [], set()
    for line_no, row in enumerate(rows[1:], start=2):
        if not any(c not in (None, "") for c in row):
            continue  # blank row
        data = {field: row[i] for i, field in mapping.items() if i < len(row)}
        try:
            entry = WatchlistEntry(**{k: v for k, v in data.items() if v is not None})
        except ValidationError as e:
            msgs = "; ".join(f"{'.'.join(map(str, er['loc']))}: {er['msg']}" for er in e.errors())
            errors.append(f"  row {line_no}: {msgs}")
            continue
        key = entry.name.lower()
        if key in seen:
            errors.append(f"  row {line_no}: duplicate company {entry.name!r}")
            continue
        seen.add(key)
        entries.append(entry)

    if errors:
        raise ConfigError(f"Problems in {path}:\n" + "\n".join(errors))
    return entries


# ---------------------------------------------------------------------------
# Importing into the database
# ---------------------------------------------------------------------------

_SYNCED_FIELDS = [
    "segment", "tier", "why_fit", "roles_to_watch", "warm_angle",
    "status", "local_flag", "ats_hint", "notes",
]


@dataclass
class ImportResult:
    added: list[str] = field(default_factory=list)
    updated: list[str] = field(default_factory=list)
    unchanged: int = 0
    deactivated: list[str] = field(default_factory=list)
    override_changed: list[str] = field(default_factory=list)


def import_watchlist(conn: sqlite3.Connection, entries: list[WatchlistEntry]) -> ImportResult:
    """Upsert entries into companies. Detection results are kept unless the
    board URL override changed, in which case the company is queued for re-detection."""
    result = ImportResult()
    existing = {
        r["name"].lower(): r for r in conn.execute("SELECT * FROM companies").fetchall()
    }
    incoming = set()

    conn.execute("BEGIN")
    try:
        for e in entries:
            key = e.name.lower()
            incoming.add(key)
            row = existing.get(key)
            values = {f: getattr(e, f) for f in _SYNCED_FIELDS}

            if row is None:
                cols = ["name", "board_url", *values]
                marks = ", ".join("?" * len(cols))
                conn.execute(
                    f"INSERT INTO companies ({', '.join(cols)}) VALUES ({marks})",
                    [e.name, e.board_url, *values.values()],
                )
                result.added.append(e.name)
                continue

            changes = {f: v for f, v in values.items() if row[f] != v}
            if row["active"] == 0:
                changes["active"] = 1
            if row["board_url"] != e.board_url:
                changes["board_url"] = e.board_url
                # Override changed: force re-detection next run
                changes["detect_status"] = None
                changes["detected_at"] = None
                result.override_changed.append(e.name)
            if row["name"] != e.name:
                changes["name"] = e.name  # capitalization fix

            if changes:
                sets = ", ".join(f"{c} = ?" for c in changes)
                conn.execute(
                    f"UPDATE companies SET {sets}, updated_at = datetime('now') WHERE id = ?",
                    [*changes.values(), row["id"]],
                )
                result.updated.append(e.name)
            else:
                result.unchanged += 1

        for key, row in existing.items():
            if key not in incoming and row["active"] == 1:
                conn.execute(
                    "UPDATE companies SET active = 0, updated_at = datetime('now') WHERE id = ?",
                    (row["id"],),
                )
                result.deactivated.append(row["name"])
        conn.execute("COMMIT")
    except Exception:
        conn.execute("ROLLBACK")
        raise
    return result
