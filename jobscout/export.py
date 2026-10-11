"""Excel export of collected postings: one sortable, filterable sheet plus a per-company summary."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

NAVY = "1F2A44"
FONT = "Arial"


def title_matches(title: str, filters) -> bool:
    """The rough title lens from filters.yaml (include any keyword, exclude all excluded ones)."""
    if filters is None:
        return True
    t = title.lower()
    inc = [k.lower() for k in filters.titles.include_keywords]
    exc = [k.lower() for k in filters.titles.exclude_keywords]
    return (not inc or any(k in t for k in inc)) and not any(k in t for k in exc)


COLUMNS = [
    ("Tier", 6), ("Company", 24), ("Segment", 24), ("Title", 52), ("Location", 26),
    ("Remote", 9), ("Pay min ($)", 12), ("Pay max ($)", 12), ("Pay source", 11),
    ("Posted", 11), ("First seen", 11), ("New", 6), ("Title match", 11), ("Closed", 11),
    ("Link", 60),
]


def write_jobs_xlsx(rows, path: Path, new_ids: set[int], filters) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Jobs"
    hdr_font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
    hdr_fill = PatternFill("solid", fgColor=NAVY)
    body = Font(name=FONT, size=10)
    link = Font(name=FONT, size=10, color="1F4E99", underline="single")

    ws.append([c for c, _ in COLUMNS])
    for i, (_, w) in enumerate(COLUMNS, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        c = ws.cell(1, i)
        c.font, c.fill = hdr_font, hdr_fill
        c.alignment = Alignment(vertical="center", wrap_text=True)

    for r in rows:
        remote = {1: "Yes", 0: "No"}.get(r["remote"], "")
        ws.append([
            r["tier"], r["company"], r["segment"], r["title"], r["location"], remote,
            r["comp_min"], r["comp_max"], r["comp_source"],
            (r["posted_at"] or "")[:10] or None, (r["first_seen_at"] or "")[:10] or None,
            "Yes" if r["id"] in new_ids else "",
            "Yes" if title_matches(r["title"], filters) else "",
            (r["closed_at"] or "")[:10] or None, r["url"],
        ])
        n = ws.max_row
        for c in ws[n]:
            c.font = body
        for col in (7, 8):
            ws.cell(n, col).number_format = "$#,##0"
        if r["url"]:
            ws.cell(n, 4).hyperlink = r["url"]
            ws.cell(n, 4).font = link
            ws.cell(n, 15).hyperlink = r["url"]
            ws.cell(n, 15).font = link

    last = ws.max_row
    table = Table(displayName="Jobs", ref=f"A1:O{last}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "E2"

    # Per-company summary, computed by formulas off the Jobs sheet so it follows any edits
    sm = wb.create_sheet("By company")
    sm.append(["Tier", "Company", "Postings", "New", "Title match", "Remote", "Avg pay min ($)"])
    for i, w in enumerate([6, 28, 10, 8, 12, 9, 18], 1):
        sm.column_dimensions[sm.cell(1, i).column_letter].width = w
        c = sm.cell(1, i)
        c.font, c.fill = hdr_font, hdr_fill
    companies = sorted({(r["tier"], r["company"]) for r in rows})
    for i, (tier, name) in enumerate(companies, start=2):
        sm.cell(i, 1, tier)
        sm.cell(i, 2, name)
        sm.cell(i, 3, f"=COUNTIF(Jobs!$B$2:$B${last},B{i})")
        sm.cell(i, 4, f'=COUNTIFS(Jobs!$B$2:$B${last},B{i},Jobs!$L$2:$L${last},"Yes")')
        sm.cell(i, 5, f'=COUNTIFS(Jobs!$B$2:$B${last},B{i},Jobs!$M$2:$M${last},"Yes")')
        sm.cell(i, 6, f'=COUNTIFS(Jobs!$B$2:$B${last},B{i},Jobs!$F$2:$F${last},"Yes")')
        sm.cell(i, 7, f'=IFERROR(AVERAGEIFS(Jobs!$G$2:$G${last},Jobs!$B$2:$B${last},B{i}),"")')
        sm.cell(i, 7).number_format = "$#,##0"
        for c in sm[i]:
            c.font = body
    if companies:
        t2 = Table(displayName="ByCompany", ref=f"A1:G{len(companies) + 1}")
        t2.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
        sm.add_table(t2)
    sm.freeze_panes = "C2"
    wb.save(path)
    return path
